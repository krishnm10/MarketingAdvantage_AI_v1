# =============================================
# docx_parser_v2.py — Enhanced Word Document Parser (Production-Ready)
# Fully compatible with ingestion_v2 architecture
# Now includes unified LLM normalization toggle (global + local)
# Supports embedded image extraction via DocumentVisualInterceptorV1
# =============================================

from typing import Dict, Any, List
from docx import Document
import os
import asyncio
import re as _re

from app.utils.text_cleaner_v2 import clean_text
from app.utils.logger import log_info, log_warning
from app.services.ingestion.llm_rewriter import rewrite_batch  # ✅ LLM integration
from app.config.ingestion_settings import ENABLE_LLM_NORMALIZATION  # ✅ Global toggle

# -------------------------------------------------------------------
# Local parser-level LLM toggle
# -------------------------------------------------------------------
# True  → Force enable LLM normalization for DOCX parser
# False → Force disable LLM normalization
# None  → Inherit from global ENABLE_LLM_NORMALIZATION
LOCAL_LLM_TOGGLE = None

# -------------------------------------------------------------------
# Visual interception toggle (embedded images in DOCX)
# -------------------------------------------------------------------
LOCAL_VISUAL_INTERCEPT_TOGGLE = True
VISUAL_INTERCEPT_TIMEOUT_SEC = 60
MAX_VISUAL_EXPLANATIONS = 24


def is_llm_enabled() -> bool:
    """Determine whether LLM normalization is enabled for this parser."""
    return ENABLE_LLM_NORMALIZATION if LOCAL_LLM_TOGGLE is None else LOCAL_LLM_TOGGLE


# -------------------------------------------------------------------
# CID CHARACTER RESOLUTION (shared with pdf_parser_v2)
# -------------------------------------------------------------------
_CID_UNICODE_MAP = {
    133: "\u2026", 145: "\u2018", 146: "\u2019", 147: "\u201C",
    148: "\u201D", 149: "\u2022", 150: "\u2013", 151: "\u2014",
    160: "\u00A0", 169: "\u00A9", 174: "\u00AE", 176: "\u00B0",
    188: "\u00BC", 189: "\u00BD", 190: "\u00BE",
    210: "\u2013", 211: "\u2014", 212: "\u201C", 213: "\u201D",
}
_CID_PATTERN = _re.compile(r"\(cid:(\d+)\)")


def resolve_cid_characters(text: str) -> str:
    """Replace (cid:NNN) placeholders with their Unicode equivalents."""
    def _replace(match):
        cid = int(match.group(1))
        return _CID_UNICODE_MAP.get(cid, "\uFFFD")
    return _CID_PATTERN.sub(_replace, text)


# -------------------------------------------------------------------
# PAGE HEADER / FOOTER STRIPPING
# -------------------------------------------------------------------
_HEADER_FOOTER_PATTERNS = [
    _re.compile(r"^\s*.*?Page\s+\d+\s+of\s+\d+\s*$", _re.MULTILINE | _re.IGNORECASE),
    _re.compile(r"^\s*[\-\u2013\u2014]+\s*\d+\s*[\-\u2013\u2014]+\s*$", _re.MULTILINE),
]


def strip_page_headers_footers(text: str) -> str:
    """Remove common page header/footer lines from extracted text."""
    for pattern in _HEADER_FOOTER_PATTERNS:
        text = pattern.sub("", text)
    text = _re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# -------------------------------------------------------------------
# TEXT EXTRACTORS
# -------------------------------------------------------------------
def extract_paragraphs(doc: Document) -> List[str]:
    """Extract visible paragraph text blocks."""
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def extract_tables(doc: Document) -> List[str]:
    """Extract readable rows from DOCX tables."""
    rows_text = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                rows_text.append(" | ".join(row_cells))
    return rows_text


def extract_headers_footers(doc: Document) -> List[str]:
    """Extract header/footer text from all sections (if any)."""
    headers, footers = [], []

    for section in doc.sections:
        if hasattr(section, "header") and section.header and section.header.paragraphs:
            for p in section.header.paragraphs:
                if p.text.strip():
                    headers.append(p.text.strip())

        if hasattr(section, "footer") and section.footer and section.footer.paragraphs:
            for p in section.footer.paragraphs:
                if p.text.strip():
                    footers.append(p.text.strip())

    return headers + footers


# -------------------------------------------------------------------
# MERGER
# -------------------------------------------------------------------
def merge_blocks(blocks: List[str]) -> str:
    """Combine extracted sections into one clean text blob."""
    return "\n\n".join(blocks)


# -------------------------------------------------------------------
# MAIN PARSER PIPELINE
# -------------------------------------------------------------------
async def parse_docx(file_path: str) -> Dict[str, Any]:
    """
    High-quality DOCX ingestion parser (async-safe).
    Steps:
      1. Extract paragraphs, tables, headers/footers
      2. Merge and clean all content
      3. (Optional) Normalize text with LLM
      4. Return ingestion_v2-ready dict
    """

    log_info(f"[docx_parser_v2] Reading DOCX: {file_path}")

    # <<< PATCH: load DOCX in a background thread to avoid blocking event loop >>>
    try:
        loop = asyncio.get_running_loop()
        doc = await loop.run_in_executor(None, lambda: Document(file_path))
    except Exception as e:
        raise ValueError(f"[docx_parser_v2] Failed to read DOCX: {file_path}: {e}")

    paragraphs = extract_paragraphs(doc)
    tables = extract_tables(doc)
    headers_footers = extract_headers_footers(doc)

    all_blocks = paragraphs + tables + headers_footers

    if not all_blocks:
        raise ValueError(f"[docx_parser_v2] No readable content in file: {file_path}")

    combined = merge_blocks(all_blocks)

    # ----------------------------------------------------------------
    # CID character resolution + header/footer stripping
    # ----------------------------------------------------------------
    combined = resolve_cid_characters(combined)
    combined = strip_page_headers_footers(combined)

    # ----------------------------------------------------------------
    # Optional visual interception (embedded images, fail-open)
    # ----------------------------------------------------------------
    visual_texts: List[str] = []
    if LOCAL_VISUAL_INTERCEPT_TOGGLE:
        try:
            from app.services.ingestion.media.document_visual_interceptor_v1 import (
                DocumentVisualInterceptorV1,
            )

            interceptor = DocumentVisualInterceptorV1()
            visual_texts = await asyncio.wait_for(
                interceptor.intercept_explanations_only(
                    file_path=file_path,
                    parsed_output={},
                    file_type="docx",
                    max_visuals=MAX_VISUAL_EXPLANATIONS,
                ),
                timeout=VISUAL_INTERCEPT_TIMEOUT_SEC,
            )
            if visual_texts:
                log_info(
                    f"[docx_parser_v2] Extracted {len(visual_texts)} visual explanations"
                )
        except asyncio.TimeoutError:
            log_warning(
                f"[docx_parser_v2] Visual interception timed out after "
                f"{VISUAL_INTERCEPT_TIMEOUT_SEC}s (non-fatal)"
            )
        except Exception as e:
            log_warning(
                f"[docx_parser_v2] Visual interception failed (non-fatal): {e}"
            )

    # Merge visual explanations into combined text
    if visual_texts:
        combined = combined + "\n\n" + "\n\n".join(visual_texts)

    cleaned = clean_text(combined)
    normalized_text = cleaned

    # ----------------------------------------------------------------
    # ✅ Optional LLM normalization
    # ----------------------------------------------------------------
    if is_llm_enabled():
        try:
            log_info("[docx_parser_v2] Sending text for LLM normalization...")
            normalized_results = await rewrite_batch([cleaned])
            if normalized_results:
                normalized_text = normalized_results[0]
                log_info("[docx_parser_v2] ✅ LLM normalization complete.")
        except Exception as e:
            log_warning(f"[docx_parser_v2] ⚠️ LLM normalization failed: {e}")
    else:
        log_info("[docx_parser_v2] LLM normalization skipped (disabled).")

    # ----------------------------------------------------------------
    # ✅ Structured Output
    # ----------------------------------------------------------------
    return {
        "raw_text": combined,
        "cleaned_text": cleaned,
        "normalized_text": normalized_text,
        "blocks": len(all_blocks),
        "visual_explanations": visual_texts,
        "source_type": "docx",
        "metadata": {
            "file_name": os.path.basename(file_path),
            "blocks": len(all_blocks),
            "paragraphs": len(paragraphs),
            "table_rows": len(tables),
            "headers_footers": len(headers_footers),
            "parser": "docx_v2 (python-docx + LLM optional)",
            "llm_normalization": is_llm_enabled(),
            "visual_explanations_count": len(visual_texts),
        },
    }
